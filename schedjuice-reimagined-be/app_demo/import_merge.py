from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from django.db.models import Q
from openpyxl import load_workbook
from tenant_schemas.utils import schema_context

from app_auth.import_commit import (
    commit_import_rows,
    dedupe_prepared_rows,
    validate_import_rows,
)
from app_auth.models import User
from app_course.models import Course
from app_demo.config import ResolvedDemoConfig
from app_demo.paths import BASE_DIR
from app_demo.validation import load_yaml

_SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
_MAPPED_FIELDS = {"email", "name", "phone_number", "date_of_birth", "course_title"}
_AUTO_HEADER_MAP = {
    "email": "email",
    "email_address": "email",
    "student_email": "email",
    "name": "name",
    "student_name": "name",
    "full_name": "name",
    "phone": "phone_number",
    "phone_number": "phone_number",
    "mobile": "phone_number",
    "mobile_number": "phone_number",
    "date_of_birth": "date_of_birth",
    "dob": "date_of_birth",
    "birth_date": "date_of_birth",
    "class": "course_title",
    "course": "course_title",
    "course_title": "course_title",
}
_FICTION_USER_PREFIXES = (
    "demo-",
    "demo_",
    "fiction-",
    "fiction_",
    "sample-",
    "sample_",
    "test-",
    "test_",
)
_FICTION_COURSE_PREFIXES = (
    "demo ",
    "fiction ",
    "fiction-",
    "sample ",
    "sample-",
    "test ",
    "test-",
)


def _normalize_header(value: str) -> str:
    text = re.sub(r"[^a-z0-9]+", "_", value.strip().lower())
    return text.strip("_")


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return value


def _is_blank_row(row: list[Any]) -> bool:
    for value in row:
        if _normalize_cell(value) is not None:
            return False
    return True


def _read_csv_rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            headers = [str(v).strip() if v is not None else "" for v in next(reader)]
        except StopIteration:
            return [], []

        rows: list[list[Any]] = []
        width = len(headers)
        for raw in reader:
            row = [raw[index] if index < len(raw) else None for index in range(width)]
            if _is_blank_row(row):
                continue
            rows.append([_normalize_cell(value) for value in row])
    return headers, rows


def _read_xlsx_rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            return [], []

        headers = [str(v).strip() if v is not None else "" for v in header_row]
        width = len(headers)
        rows: list[list[Any]] = []
        for raw in iterator:
            row = [raw[index] if index < len(raw) else None for index in range(width)]
            if _is_blank_row(row):
                continue
            rows.append([_normalize_cell(value) for value in row])
        return headers, rows
    finally:
        workbook.close()


def _read_sheet_rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    extension = path.suffix.lower()
    if extension not in _SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(_SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported import file extension '{extension}'. Use {supported}.")
    if extension == ".csv":
        return _read_csv_rows(path)
    return _read_xlsx_rows(path)


def _resolve_import_path(file_ref: str) -> Path:
    candidate = Path(file_ref)
    if not candidate.is_absolute():
        candidate = (BASE_DIR / candidate).resolve()
    if not candidate.exists():
        raise ValueError(f"Import file not found: {candidate}")
    return candidate


def _header_lookup(headers: list[str]) -> dict[str, str]:
    return {header.strip().lower(): header for header in headers if header.strip()}


def _resolve_column_mapping(
    *,
    config: ResolvedDemoConfig,
    headers: list[str],
    sidecar_columns: dict[str, Any] | None,
) -> dict[str, str]:
    if sidecar_columns is not None:
        mapped: dict[str, str] = {}
        used_targets: dict[str, str] = {}
        lookup = _header_lookup(headers)
        for raw_header, raw_target in sidecar_columns.items():
            if not isinstance(raw_header, str) or not str(raw_header).strip():
                continue
            if not isinstance(raw_target, str) or raw_target.strip() not in _MAPPED_FIELDS:
                raise ValueError(
                    f"Invalid mapping target for '{raw_header}'. "
                    f"Allowed targets: {', '.join(sorted(_MAPPED_FIELDS))}."
                )
            source_header = lookup.get(raw_header.strip().lower())
            if source_header is None:
                raise ValueError(f"Mapping references unknown column '{raw_header}'.")
            target = raw_target.strip()
            if target in used_targets and used_targets[target] != source_header:
                raise ValueError(
                    f"Ambiguous mapping: both '{used_targets[target]}' and "
                    f"'{source_header}' map to '{target}'."
                )
            mapped[source_header] = target
            used_targets[target] = source_header
        if not mapped:
            raise ValueError("Mapping sidecar has no usable column mappings.")
        return mapped

    mapped = {}
    unmapped: list[str] = []
    used_targets: dict[str, str] = {}
    ambiguous_pairs: list[tuple[str, str, str]] = []
    for header in headers:
        key = _normalize_header(header)
        target = _AUTO_HEADER_MAP.get(key)
        if target is None:
            unmapped.append(header)
            continue
        if target in used_targets:
            ambiguous_pairs.append((used_targets[target], header, target))
            continue
        mapped[header] = target
        used_targets[target] = header

    if unmapped or ambiguous_pairs:
        template_path = _write_mapping_template(config=config, headers=headers)
        pieces: list[str] = []
        if unmapped:
            pieces.append(f"Unmapped headers: {', '.join(unmapped)}.")
        for first, second, target in ambiguous_pairs:
            pieces.append(f"Ambiguous columns '{first}' and '{second}' both map to '{target}'.")
        pieces.append(f"Fill mapping file: {template_path}")
        raise ValueError(" ".join(pieces))

    return mapped


def _split_course_titles(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        values = value
    else:
        values = [value]

    titles: list[str] = []
    for item in values:
        if item is None:
            continue
        text = str(item).strip()
        if not text:
            continue
        parts = re.split(r"[;,|]", text)
        for part in parts:
            normalized = part.strip()
            if normalized:
                titles.append(normalized)
    return titles


def _write_mapping_template(*, config: ResolvedDemoConfig, headers: list[str]) -> Path:
    template_path = (
        BASE_DIR / "demo-artifacts" / "imports" / f"{config.slug}.mapping.template.yaml"
    )
    template_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "role: student",
        "columns:",
    ]
    for header in headers:
        escaped = header.replace('"', '\\"')
        lines.append(f'  "{escaped}": ""')
    template_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return template_path


def _format_errors(errors: list[dict[str, Any]]) -> str:
    lines = ["Import validation failed:"]
    for error in errors[:50]:
        row = int(error.get("row", 0)) + 2
        field = error.get("field", "row")
        reason = error.get("reason", "Unknown error.")
        lines.append(f"- row {row} [{field}]: {reason}")
    if len(errors) > 50:
        lines.append(f"- ... {len(errors) - 50} more error(s)")
    return "\n".join(lines)


def _replace_fiction_rows() -> None:
    user_filter = Q()
    for prefix in _FICTION_USER_PREFIXES:
        user_filter |= Q(email__istartswith=prefix)
    User.objects.filter(user_filter).delete()

    course_filter = Q()
    for prefix in _FICTION_COURSE_PREFIXES:
        course_filter |= Q(title__istartswith=prefix)
    Course.objects.exclude(title__startswith="DEMO-").filter(course_filter).delete()


def _resolve_course_ids(course_titles: set[str]) -> tuple[dict[str, int], set[str]]:
    if not course_titles:
        return {}, set()
    matched = {
        title: course_id
        for title, course_id in Course.objects.filter(title__in=course_titles).values_list(
            "title", "id"
        )
    }
    missing = set(course_titles) - set(matched.keys())
    return matched, missing


def _load_sidecar(path: Path) -> tuple[str, dict[str, Any] | None]:
    sidecar_path = path.with_suffix(".mapping.yaml")
    if not sidecar_path.exists():
        return User.UserRole.STUDENT, None

    sidecar = load_yaml(sidecar_path)
    role = sidecar.get("role") or User.UserRole.STUDENT
    if role not in User.UserRole.values:
        supported = ", ".join(User.UserRole.values)
        raise ValueError(f"Invalid import role '{role}'. Use one of: {supported}.")

    columns = sidecar.get("columns")
    if columns is not None and not isinstance(columns, dict):
        raise ValueError(f"Invalid mapping file at {sidecar_path}: 'columns' must be a mapping.")
    return role, columns


def run_import_merge(
    *,
    config: ResolvedDemoConfig,
    schema_name: str,
    structure_ctx: dict,
    pack_ctx: dict,
    skip: bool = False,
) -> dict:
    if skip or not config.import_spec:
        return {"skipped": True}

    del structure_ctx, pack_ctx

    import_file = str(config.import_spec.get("file") or "").strip()
    if not import_file:
        raise ValueError("Brief import.file is required when import block is present.")

    mode = str(config.import_spec.get("mode") or "merge").strip().lower()
    if mode not in {"merge", "replace_fiction"}:
        raise ValueError("Brief import.mode must be one of: merge, replace_fiction.")

    file_path = _resolve_import_path(import_file)
    headers, raw_rows = _read_sheet_rows(file_path)
    if not headers:
        raise ValueError(f"Import file has no header row: {file_path}")
    if not raw_rows:
        return {"skipped": False, "imported_count": 0, "errors": []}

    role, sidecar_columns = _load_sidecar(file_path)
    column_mapping = _resolve_column_mapping(
        config=config,
        headers=headers,
        sidecar_columns=sidecar_columns,
    )
    if "email" not in set(column_mapping.values()):
        raise ValueError("Import mapping must include an email column.")

    mapped_rows: list[dict[str, Any]] = []
    course_titles_by_row: list[list[str]] = []
    all_course_titles: set[str] = set()
    for raw in raw_rows:
        mapped: dict[str, Any] = {"custom_data": {}, "course_ids": []}
        course_titles: list[str] = []
        for index, header in enumerate(headers):
            target = column_mapping.get(header)
            if target is None:
                continue
            value = raw[index] if index < len(raw) else None
            if target == "course_title":
                titles = _split_course_titles(value)
                course_titles.extend(titles)
                all_course_titles.update(titles)
                continue
            mapped[target] = value
        mapped_rows.append(mapped)
        course_titles_by_row.append(course_titles)

    with schema_context(schema_name):
        if mode == "replace_fiction":
            _replace_fiction_rows()

        course_lookup, missing_titles = _resolve_course_ids(all_course_titles)
        errors: list[dict[str, Any]] = []
        if missing_titles:
            missing_display = ", ".join(sorted(missing_titles))
            for index, row_titles in enumerate(course_titles_by_row):
                row_missing = [title for title in row_titles if title in missing_titles]
                if row_missing:
                    errors.append(
                        {
                            "row": index,
                            "field": "course_title",
                            "reason": (
                                "Unknown course title(s): "
                                f"{', '.join(row_missing)}. Available titles do not include these."
                            ),
                        }
                    )
            if not errors:
                errors.append(
                    {
                        "row": 0,
                        "field": "course_title",
                        "reason": f"Unknown course title(s): {missing_display}",
                    }
                )

        for index, mapped_row in enumerate(mapped_rows):
            row_titles = course_titles_by_row[index]
            unique_course_ids: list[int] = []
            seen_ids: set[int] = set()
            for title in row_titles:
                course_id = course_lookup.get(title)
                if course_id is None or course_id in seen_ids:
                    continue
                seen_ids.add(course_id)
                unique_course_ids.append(course_id)
            mapped_row["course_ids"] = unique_course_ids

        validate_errors, prepared = validate_import_rows(rows=mapped_rows, role=role)
        if validate_errors:
            errors.extend(validate_errors)
        if errors:
            raise ValueError(_format_errors(errors))

        prepared = dedupe_prepared_rows(prepared, strategy="merge")
        imported_count = len(prepared)
        commit_import_rows(prepared=prepared, role=role)

    return {
        "skipped": False,
        "imported_count": imported_count,
        "errors": [],
    }
