"""Stateless spreadsheet parsing for the Import Wizard."""
from __future__ import annotations

import csv
from typing import BinaryIO

from openpyxl import load_workbook

CSV_SHEET_NAME = "CSV"


def _cell(value) -> str | int | float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    return str(value).strip()


def parse_workbook(file_obj: BinaryIO, *, sheet: str | None, max_rows: int) -> dict:
    wb = load_workbook(file_obj, data_only=True, read_only=True)
    sheet_names = list(wb.sheetnames)
    active = sheet if sheet in sheet_names else sheet_names[0]
    ws = wb[active]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        return {
            "sheet_names": sheet_names,
            "active_sheet": active,
            "headers": [],
            "rows": [],
            "row_count": 0,
        }

    headers = [("" if h is None else str(h)) for h in header_row]
    width = len(headers)

    rows: list[list] = []
    for raw in rows_iter:
        if raw is None or all(c is None for c in raw):
            continue
        cells = [_cell(raw[i]) if i < len(raw) else None for i in range(width)]
        rows.append(cells)
        if len(rows) > max_rows:
            wb.close()
            raise ValueError(f"File exceeds the maximum of {max_rows} rows.")

    wb.close()
    return {
        "sheet_names": sheet_names,
        "active_sheet": active,
        "headers": headers,
        "rows": rows,
        "row_count": len(rows),
    }


def parse_csv(file_obj: BinaryIO, *, max_rows: int, delimiter: str = ",") -> dict:
    text = file_obj.read().decode("utf-8-sig")
    reader = csv.reader(text.splitlines(), delimiter=delimiter)

    try:
        header_row = next(reader)
    except StopIteration:
        return {
            "sheet_names": [CSV_SHEET_NAME],
            "active_sheet": CSV_SHEET_NAME,
            "headers": [],
            "rows": [],
            "row_count": 0,
        }

    headers = [("" if h is None else str(h)) for h in header_row]
    width = len(headers)

    rows: list[list] = []
    for raw in reader:
        if not raw or all(c is None or str(c).strip() == "" for c in raw):
            continue
        cells = [_cell(raw[i]) if i < len(raw) else None for i in range(width)]
        rows.append(cells)
        if len(rows) > max_rows:
            raise ValueError(f"File exceeds the maximum of {max_rows} rows.")

    return {
        "sheet_names": [CSV_SHEET_NAME],
        "active_sheet": CSV_SHEET_NAME,
        "headers": headers,
        "rows": rows,
        "row_count": len(rows),
    }


def parse_tsv(file_obj: BinaryIO, *, max_rows: int) -> dict:
    return parse_csv(file_obj, max_rows=max_rows, delimiter="\t")


def parse_import_file(
    file_obj: BinaryIO,
    *,
    filename: str | None,
    sheet: str | None,
    max_rows: int,
) -> dict:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return parse_csv(file_obj, max_rows=max_rows)
    if name.endswith(".tsv"):
        return parse_tsv(file_obj, max_rows=max_rows)
    return parse_workbook(file_obj, sheet=sheet, max_rows=max_rows)
