import io
import unittest
from datetime import date
from uuid import uuid4

from django.core.management import call_command
from django.db import connection
from django.test import SimpleTestCase, TestCase, override_settings
from openpyxl import Workbook
from tenant_schemas.utils import schema_context

from app_auth.models import User
from app_course.models import Course
from app_grading_reports.mark_sheet_header_flatten import flatten_two_row_headers
from app_grading_reports.mark_sheet_services import infer_rubric_columns
from app_rbac.seeding import seed_rbac


def _database_reachable() -> bool:
    try:
        connection.ensure_connection()
        return True
    except Exception:
        return False


SCREENSHOT_HEADERS = [
    "No.",
    "Burmese Name",
    "English name",
    "Reading and Use of English (78 marks)",
    "",
    "Writing (40 marks)",
    "",
    "Listening (30 marks)",
    "",
    "Speaking (25 marks)",
    "",
]

SCREENSHOT_SUB_ROW = [
    "",
    "",
    "",
    "Mark",
    "Grade",
    "Mark",
    "Grade",
    "Mark",
    "Grade",
    "Mark",
    "Grade",
]

SCREENSHOT_DATA_ROW = [
    "1",
    "Aung Myin Su Naing",
    "Paul Herbold",
    "53",
    "B",
    "28",
    "B",
    "29",
    "A+",
    "22",
    "A+",
]


class MarkSheetHeaderFlattenTests(SimpleTestCase):
    def test_flattens_two_row_merged_header_layout(self):
        headers, rows = flatten_two_row_headers(
            SCREENSHOT_HEADERS,
            [SCREENSHOT_SUB_ROW, SCREENSHOT_DATA_ROW],
        )

        self.assertIn("Reading and Use of English (78 marks) Mark", headers)
        self.assertIn("Reading and Use of English (78 marks) Grade", headers)
        self.assertIn("Writing (40 marks) Mark", headers)
        self.assertEqual(headers[0], "No.")
        self.assertEqual(headers[2], "English name")
        self.assertEqual(rows[0], SCREENSHOT_DATA_ROW)

    def test_single_row_header_unchanged(self):
        headers = ["No.", "English name", "Content", "Total (40) mark"]
        rows = [["1", "Paul", "4", "28"]]

        flat_headers, flat_rows = flatten_two_row_headers(headers, rows)

        self.assertEqual(flat_headers, headers)
        self.assertEqual(flat_rows, rows)

    def test_does_not_flatten_when_sub_row_looks_like_student_data(self):
        headers = SCREENSHOT_HEADERS
        rows = [SCREENSHOT_DATA_ROW, ["2", "Other", "David", "50", "B", "30", "B", "28", "A", "20", "A"]]

        flat_headers, flat_rows = flatten_two_row_headers(headers, rows)

        self.assertEqual(flat_headers, headers)
        self.assertEqual(flat_rows, rows)

    def test_mark_columns_score_and_grade_columns_identifier(self):
        headers, rows = flatten_two_row_headers(
            SCREENSHOT_HEADERS,
            [SCREENSHOT_SUB_ROW, SCREENSHOT_DATA_ROW],
        )
        cols = infer_rubric_columns(headers, rows)
        kinds = {c["title"]: c["kind"] for c in cols}

        self.assertEqual(kinds["Reading and Use of English (78 marks) Mark"], "score")
        self.assertEqual(kinds["Writing (40 marks) Mark"], "score")
        self.assertEqual(kinds["Reading and Use of English (78 marks) Grade"], "ignored")
        self.assertEqual(kinds["English name"], "identifier")


def _merged_mark_sheet_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Marks"
    ws.append(SCREENSHOT_HEADERS)
    ws.append(SCREENSHOT_SUB_ROW)
    ws.append(SCREENSHOT_DATA_ROW)
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:C2")
    ws.merge_cells("D1:E1")
    ws.merge_cells("F1:G1")
    ws.merge_cells("H1:I1")
    ws.merge_cells("J1:K1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@unittest.skipUnless(_database_reachable(), "PostgreSQL not available")
@override_settings(RBAC_ENFORCE="enforce")
class MarkSheetImportParseViewTests(TestCase):
    schema_name = "xschedjuice"
    api_prefix = "/api/v1"

    @classmethod
    def setUpTestData(cls):
        call_command("migrate_schemas", shared=True, verbosity=0)
        call_command("migrate_schemas", verbosity=0)
        call_command("load-data", schema=cls.schema_name, verbosity=0)

    def setUp(self):
        self.suffix = uuid4().hex[:6]
        with schema_context(self.schema_name):
            seed_rbac()
            self.course = Course.objects.first()
            self.admin = User.objects.create_user(
                email=f"msh-admin-{self.suffix}@example.com",
                password="x",
                name="Ms Admin",
                phone_number="1",
                date_of_birth=date(1990, 1, 1),
                communication_email=f"msh-admin-{self.suffix}@example.com",
                code=f"msh-admin-{self.suffix}",
                roles=[User.UserRole.ADMIN],
            )

    def _client(self, user):
        from rest_framework.test import APIClient

        client = APIClient()
        client.force_authenticate(user=user)
        client.credentials(HTTP_TENANT=self.schema_name)
        return client

    def test_parse_merged_xlsx_flattens_headers(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile(
            "marks.xlsx",
            _merged_mark_sheet_xlsx(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        with schema_context(self.schema_name):
            res = self._client(self.admin).post(
                f"{self.api_prefix}/courses/{self.course.id}/mark-sheets/import/parse",
                {"file": upload},
                format="multipart",
            )
        self.assertEqual(res.status_code, 200, res.content)
        data = res.json()["data"]
        self.assertIn("Reading and Use of English (78 marks) Mark", data["headers"])
        self.assertEqual(data["rows"][0][2], "Paul Herbold")
        score_titles = [c["title"] for c in data["inferred_columns"] if c["kind"] == "score"]
        self.assertIn("Reading and Use of English (78 marks) Mark", score_titles)

    def test_parse_paste_flattens_two_row_header(self):
        paste_lines = [
            "\t".join(SCREENSHOT_HEADERS),
            "\t".join(SCREENSHOT_SUB_ROW),
            "\t".join(SCREENSHOT_DATA_ROW),
        ]
        with schema_context(self.schema_name):
            res = self._client(self.admin).post(
                f"{self.api_prefix}/courses/{self.course.id}/mark-sheets/import/parse",
                {"paste": "\n".join(paste_lines)},
                format="json",
            )
        self.assertEqual(res.status_code, 200, res.content)
        data = res.json()["data"]
        self.assertIn("Writing (40 marks) Mark", data["headers"])
        self.assertEqual(data["rows"][0][0], "1")
