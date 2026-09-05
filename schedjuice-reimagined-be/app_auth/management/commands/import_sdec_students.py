"""
Import SDEC students from per-room workbook sheets into a tenant schema.

Each sheet (e.g. Y3-R1) maps to an existing course. Resolves email and
microsoft_id from a Microsoft directory CSV by exact normalized displayName match.

Usage:
  ./env/bin/python manage.py import_sdec_students \\
    --xlsx=/path/sdec-student-w-room.xlsx \\
    --object-id-csv=/path/sdec-user-object-id-list.csv \\
    --dry-run

  ./env/bin/python manage.py import_sdec_students \\
    --schema-name=xsdecschedjuicecom \\
    --xlsx=/path/sdec-student-w-room.xlsx \\
    --object-id-csv=/path/sdec-user-object-id-list.csv \\
    --skip-report-csv=/path/sdec-student-w-room-skips.csv
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import make_password
from django.core.management import BaseCommand
from django.core.management.base import CommandError
from django.db import transaction
from openpyxl import load_workbook
from tenant_schemas.utils import get_public_schema_name, schema_context

from app_course.models import Course, UserCourse
from app_organization.acca_spreadsheet_import import IMPORT_PASSWORD, normalize_email
from app_organization.models import Organization

UserModel = get_user_model()

DEFAULT_SCHEMA_NAME = "xsdecschedjuicecom"
IMPORT_DATE_OF_BIRTH = date(1900, 1, 1)
PLACEHOLDER_PHONE = "-"
ACADEMIC_YEAR_SUFFIX = "Academic Year 2026-2027"
BULK_BATCH_SIZE = 500

SKIP_SHEETS = frozenset({"Nursery", "Reception 1", "Reception 2"})
SHEET_PATTERN = re.compile(r"^Y(\d+)-R(\d+)$", re.IGNORECASE)

SKIP_REPORT_COLUMNS = ("student name", "course name", "skip reason")


@dataclass
class DirectoryEntry:
    display_name: str
    email: str
    microsoft_id: str


@dataclass
class SkippedStudent:
    student_name: str
    course_name: str
    skip_reason: str


@dataclass
class PendingEnrollment:
    entry: DirectoryEntry
    course: Course


def _is_blank(value: str | None) -> bool:
    return value is None or str(value).strip() == ""


def normalize_display_name(value: str) -> str:
    """Lowercase, strip punctuation/parentheticals/@ for directory matching."""
    s = str(value).strip().lower()
    s = re.sub(r"\(.*?\)", "", s)
    s = s.replace("@", " ")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def sheet_to_course_title(sheet_name: str) -> str | None:
    match = SHEET_PATTERN.match(sheet_name.strip())
    if not match:
        return None
    year_num, section_num = match.groups()
    return (
        f"Year {int(year_num)} Section R{int(section_num)} - {ACADEMIC_YEAR_SUFFIX}"
    )


def _ambiguous_skip_reason(candidates: list[DirectoryEntry]) -> str:
    parts = [
        f"{candidate.display_name} <{candidate.email}>"
        for candidate in candidates
    ]
    joined = ", ".join(parts)
    return (
        f"ambiguous directory match ({len(candidates)} candidates: {joined})"
    )


def _write_skip_report_csv(path: Path, rows: list[SkippedStudent]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=SKIP_REPORT_COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "student name": row.student_name,
                    "course name": row.course_name,
                    "skip reason": row.skip_reason,
                }
            )


def _load_directory_index(object_id_csv: Path) -> dict[str, list[DirectoryEntry]]:
    index: dict[str, list[DirectoryEntry]] = defaultdict(list)
    with object_id_csv.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            display_name = (row.get("displayName") or "").strip()
            upn = (row.get("userPrincipalName") or "").strip()
            obj_id = (row.get("id") or "").strip()
            if _is_blank(display_name) or _is_blank(upn) or _is_blank(obj_id):
                continue
            key = normalize_display_name(display_name)
            if not key:
                continue
            index[key].append(
                DirectoryEntry(
                    display_name=display_name,
                    email=normalize_email(upn),
                    microsoft_id=obj_id,
                )
            )
    return index


def _extract_student_names_from_sheet(rows: list[tuple]) -> list[str]:
    """Extract primary (cols A/B) and secondary New Students (cols D/E) names."""
    names: list[str] = []
    for row in rows:
        if not row:
            continue
        col_a = row[0] if len(row) > 0 else None
        col_b = row[1] if len(row) > 1 else None
        if isinstance(col_a, int) and isinstance(col_b, str) and col_b.strip():
            names.append(col_b.strip())

        col_d = row[3] if len(row) > 3 else None
        col_e = row[4] if len(row) > 4 else None
        if isinstance(col_d, int) and isinstance(col_e, str) and col_e.strip():
            names.append(col_e.strip())
    return names


def _resolve_directory_entry(
    raw_name: str,
    directory_index: dict[str, list[DirectoryEntry]],
) -> tuple[DirectoryEntry | None, list[DirectoryEntry]]:
    key = normalize_display_name(raw_name)
    if not key:
        return None, []
    hits = directory_index.get(key, [])
    if len(hits) == 1:
        return hits[0], []
    if len(hits) > 1:
        return None, hits
    return None, []


def _dedupe_enrollments(
    pending: list[PendingEnrollment],
) -> list[PendingEnrollment]:
    seen: set[tuple[str, int]] = set()
    deduped: list[PendingEnrollment] = []
    for item in pending:
        key = (item.entry.email.lower(), item.course.id)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def _build_user_instance(entry: DirectoryEntry, hashed_password: str) -> UserModel:
    return UserModel(
        email=entry.email,
        password=hashed_password,
        name=entry.display_name,
        phone_number=PLACEHOLDER_PHONE,
        communication_email=entry.email,
        roles=[UserModel.UserRole.STUDENT],
        date_of_birth=IMPORT_DATE_OF_BIRTH,
        microsoft_id=entry.microsoft_id,
        is_active=True,
        is_staff=False,
        is_password_change_required=True,
    )


class Command(BaseCommand):
    help = (
        "Import SDEC student room workbook into a tenant (omit --dry-run to write). "
        "Resolves users from object-id CSV by displayName."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--schema-name",
            type=str,
            default=DEFAULT_SCHEMA_NAME,
            help=f"Tenant schema_name (default {DEFAULT_SCHEMA_NAME!r})",
        )
        parser.add_argument(
            "--xlsx",
            type=str,
            required=True,
            help="Path to sdec-student-w-room.xlsx",
        )
        parser.add_argument(
            "--object-id-csv",
            type=str,
            required=True,
            help="Path to sdec-user-object-id-list.csv (UTF-8)",
        )
        parser.add_argument(
            "--skip-report-csv",
            type=str,
            default=None,
            help=(
                "Path for skipped/ambiguous student report CSV "
                "(default: <xlsx-stem>-skips.csv beside the workbook)"
            ),
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and report only; do not write to the database",
        )

    def handle(self, *args, **options):
        schema_name: str = options["schema_name"]
        xlsx_path = Path(options["xlsx"]).expanduser().resolve()
        object_id_path = Path(options["object_id_csv"]).expanduser().resolve()
        skip_report_path = (
            Path(options["skip_report_csv"]).expanduser().resolve()
            if options["skip_report_csv"]
            else xlsx_path.parent / f"{xlsx_path.stem}-skips.csv"
        )
        dry_run: bool = options["dry_run"]
        commit: bool = not dry_run

        if not xlsx_path.is_file():
            raise CommandError(f"Workbook not found: {xlsx_path}")
        if not object_id_path.is_file():
            raise CommandError(f"Object id CSV not found: {object_id_path}")

        with schema_context(get_public_schema_name()):
            if not Organization.objects.filter(schema_name=schema_name).exists():
                raise CommandError(f"No organization for schema_name={schema_name!r}")

        if dry_run:
            self.stdout.write(
                self.style.WARNING(
                    "DRY RUN — no database writes "
                    "(counts below are planned creates/reuses; DB lookups still run)."
                )
            )

        directory_index = _load_directory_index(object_id_path)
        workbook = load_workbook(xlsx_path, data_only=True)

        stats = {
            "sheets_processed": 0,
            "sheets_skipped_no_course": 0,
            "student_rows_read": 0,
            "resolved_unique": 0,
            "skipped_no_directory_match": 0,
            "skipped_ambiguous": 0,
            "users_created": 0,
            "users_reused": 0,
            "microsoft_ids_backfilled": 0,
            "enrollments_created": 0,
            "enrollments_existing": 0,
        }
        skipped_rows: list[SkippedStudent] = []
        pending_enrollments: list[PendingEnrollment] = []

        with schema_context(schema_name):
            expected_titles = [
                title
                for ws in workbook.worksheets
                if ws.title not in SKIP_SHEETS
                for title in [sheet_to_course_title(ws.title)]
                if title is not None
            ]
            course_by_title = {
                course.title: course
                for course in Course.objects.filter(title__in=expected_titles)
            }

            for worksheet in workbook.worksheets:
                sheet_name = worksheet.title
                if sheet_name in SKIP_SHEETS:
                    continue

                course_title = sheet_to_course_title(sheet_name)
                if course_title is None:
                    self.stdout.write(
                        self.style.NOTICE(
                            f"sheet {sheet_name!r}: SKIP — unrecognised sheet name"
                        )
                    )
                    stats["sheets_skipped_no_course"] += 1
                    continue

                course = course_by_title.get(course_title)
                if course is None:
                    self.stdout.write(
                        self.style.NOTICE(
                            f"sheet {sheet_name!r}: SKIP — no course with title {course_title!r}"
                        )
                    )
                    stats["sheets_skipped_no_course"] += 1
                    continue

                stats["sheets_processed"] += 1
                rows = list(worksheet.iter_rows(values_only=True))
                student_names = _extract_student_names_from_sheet(rows)

                for raw_name in student_names:
                    stats["student_rows_read"] += 1
                    entry, ambiguous_hits = _resolve_directory_entry(
                        raw_name, directory_index
                    )

                    if ambiguous_hits:
                        stats["skipped_ambiguous"] += 1
                        skipped_rows.append(
                            SkippedStudent(
                                student_name=raw_name,
                                course_name=course.title,
                                skip_reason=_ambiguous_skip_reason(ambiguous_hits),
                            )
                        )
                        self.stdout.write(
                            self.style.NOTICE(
                                f"[{sheet_name}] SKIP — ambiguous directory match for "
                                f"{raw_name!r} ({len(ambiguous_hits)} candidates)"
                            )
                        )
                        continue

                    if entry is None:
                        stats["skipped_no_directory_match"] += 1
                        skipped_rows.append(
                            SkippedStudent(
                                student_name=raw_name,
                                course_name=course.title,
                                skip_reason="no directory match",
                            )
                        )
                        self.stdout.write(
                            self.style.NOTICE(
                                f"[{sheet_name}] SKIP — no directory match for {raw_name!r}"
                            )
                        )
                        continue

                    stats["resolved_unique"] += 1
                    pending_enrollments.append(
                        PendingEnrollment(entry=entry, course=course)
                    )

            pending_enrollments = _dedupe_enrollments(pending_enrollments)
            self._apply_enrollments(
                pending_enrollments=pending_enrollments,
                commit=commit,
                stats=stats,
            )

        self.stdout.write(self.style.SUCCESS(str(stats)))
        self._print_ambiguous_summary(skipped_rows)
        _write_skip_report_csv(skip_report_path, skipped_rows)
        self.stdout.write(
            self.style.SUCCESS(
                f"Skip report written to {skip_report_path} ({len(skipped_rows)} rows)"
            )
        )

    def _apply_enrollments(
        self,
        *,
        pending_enrollments: list[PendingEnrollment],
        commit: bool,
        stats: dict,
    ) -> None:
        if not pending_enrollments:
            return

        entries_by_email: dict[str, DirectoryEntry] = {}
        for item in pending_enrollments:
            entries_by_email.setdefault(item.entry.email.lower(), item.entry)

        emails = list(entries_by_email.keys())
        course_ids = {item.course.id for item in pending_enrollments}

        existing_users = {
            user.email.lower(): user
            for user in UserModel.objects.filter(email__in=emails)
        }
        existing_enrollment_keys = set(
            UserCourse.objects.filter(
                user__email__in=emails,
                course_id__in=course_ids,
            ).values_list("user__email", "course_id")
        )
        existing_enrollment_keys = {
            (email.lower(), course_id)
            for email, course_id in existing_enrollment_keys
        }

        seen_emails: set[str] = set(existing_users.keys())
        backfilled_emails: set[str] = set()
        planned_enrollment_keys: set[tuple[str, int]] = set()

        for item in pending_enrollments:
            email_key = item.entry.email.lower()
            user = existing_users.get(email_key)

            if user is None:
                if email_key not in seen_emails:
                    stats["users_created"] += 1
                    seen_emails.add(email_key)
                else:
                    stats["users_reused"] += 1
            else:
                stats["users_reused"] += 1
                if (
                    _is_blank(user.microsoft_id)
                    and item.entry.microsoft_id
                    and email_key not in backfilled_emails
                ):
                    stats["microsoft_ids_backfilled"] += 1
                    backfilled_emails.add(email_key)

            enrollment_key = (email_key, item.course.id)
            if (
                enrollment_key in existing_enrollment_keys
                or enrollment_key in planned_enrollment_keys
            ):
                stats["enrollments_existing"] += 1
            else:
                stats["enrollments_created"] += 1
                planned_enrollment_keys.add(enrollment_key)

        if not commit:
            return

        with transaction.atomic():
            self._bulk_create_users(
                entries_by_email=entries_by_email,
                existing_users=existing_users,
            )
            self._bulk_backfill_microsoft_ids(
                entries_by_email=entries_by_email,
                existing_users=existing_users,
            )
            user_by_email = {
                user.email.lower(): user
                for user in UserModel.objects.filter(email__in=emails)
            }
            self._bulk_create_enrollments(
                pending_enrollments=pending_enrollments,
                user_by_email=user_by_email,
                existing_enrollment_keys=existing_enrollment_keys,
            )

    def _bulk_create_users(
        self,
        *,
        entries_by_email: dict[str, DirectoryEntry],
        existing_users: dict[str, UserModel],
    ) -> None:
        emails_to_create = [
            email
            for email in entries_by_email
            if email not in existing_users
        ]
        if not emails_to_create:
            return

        hashed_password = make_password(IMPORT_PASSWORD)
        users_to_create = [
            _build_user_instance(entries_by_email[email], hashed_password)
            for email in emails_to_create
        ]
        UserModel.objects.bulk_create(users_to_create, batch_size=BULK_BATCH_SIZE)

    def _bulk_backfill_microsoft_ids(
        self,
        *,
        entries_by_email: dict[str, DirectoryEntry],
        existing_users: dict[str, UserModel],
    ) -> None:
        users_to_update: list[UserModel] = []
        for email_key, user in existing_users.items():
            entry = entries_by_email.get(email_key)
            if entry is None:
                continue
            if _is_blank(user.microsoft_id) and entry.microsoft_id:
                user.microsoft_id = entry.microsoft_id
                users_to_update.append(user)

        if users_to_update:
            UserModel.objects.bulk_update(
                users_to_update,
                ["microsoft_id"],
                batch_size=BULK_BATCH_SIZE,
            )

    def _bulk_create_enrollments(
        self,
        *,
        pending_enrollments: list[PendingEnrollment],
        user_by_email: dict[str, UserModel],
        existing_enrollment_keys: set[tuple[str, int]],
    ) -> None:
        enrollments_to_create: list[UserCourse] = []
        seen_keys = set(existing_enrollment_keys)

        for item in pending_enrollments:
            email_key = item.entry.email.lower()
            user = user_by_email.get(email_key)
            if user is None:
                continue

            enrollment_key = (email_key, item.course.id)
            if enrollment_key in seen_keys:
                continue
            seen_keys.add(enrollment_key)
            enrollments_to_create.append(
                UserCourse(
                    user_id=user.id,
                    course_id=item.course.id,
                    assigned_as=UserCourse.AssignedAs.STUDENT,
                )
            )

        if enrollments_to_create:
            UserCourse.objects.bulk_create(
                enrollments_to_create,
                batch_size=BULK_BATCH_SIZE,
                ignore_conflicts=True,
            )

    def _print_ambiguous_summary(self, skipped_rows: list[SkippedStudent]) -> None:
        ambiguous_rows = [
            row
            for row in skipped_rows
            if row.skip_reason.startswith("ambiguous directory match")
        ]
        if not ambiguous_rows:
            return

        self.stdout.write("")
        self.stdout.write(
            self.style.WARNING(
                f"AMBIGUOUS STUDENTS ({len(ambiguous_rows)}):"
            )
        )
        for item in ambiguous_rows:
            self.stdout.write(
                f"  [{item.course_name}] {item.student_name!r} -> {item.skip_reason}"
            )
