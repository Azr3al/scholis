"""
Import SDEC IGCSE batch student lists into a tenant schema (local users only).

Each workbook is a single sheet with columns No / Student Name / Username.
Resolves microsoft_id from a Microsoft directory CSV by userPrincipalName.
Does not create Microsoft accounts via Graph and does not enroll into courses.

Usage:
  ./env/bin/python manage.py import_sdec_ig_students \\
    --xlsx=/path/sdec-IG\\ Batch-19.xlsx \\
    --xlsx=/path/sdec-IG\\ Batch-20.xlsx \\
    --object-id-csv=/path/sdec-user-object-id-list.csv \\
    --dry-run

  ./env/bin/python manage.py import_sdec_ig_students \\
    --schema-name=xsdecschedjuicecom \\
    --xlsx=/path/sdec-IG\\ Batch-19.xlsx \\
    --xlsx=/path/sdec-IG\\ Batch-20.xlsx \\
    --object-id-csv=/path/sdec-user-object-id-list.csv
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import make_password
from django.core.management import BaseCommand
from django.core.management.base import CommandError
from django.db import transaction
from openpyxl import load_workbook
from tenant_schemas.utils import get_public_schema_name, schema_context

from app_organization.acca_spreadsheet_import import IMPORT_PASSWORD, normalize_email
from app_organization.models import Organization

UserModel = get_user_model()

DEFAULT_SCHEMA_NAME = "xsdecschedjuicecom"
IMPORT_DATE_OF_BIRTH = date(1900, 1, 1)
PLACEHOLDER_PHONE = "-"
BULK_BATCH_SIZE = 500

HEADER_NO = "no"
HEADER_STUDENT_NAME = "student name"
HEADER_USERNAME = "username"

SKIP_REPORT_COLUMNS = ("source file", "student name", "username", "skip reason")


@dataclass
class ParsedRow:
    source_file: str
    student_name: str
    username: str


@dataclass
class SkippedRow:
    source_file: str
    student_name: str
    username: str
    skip_reason: str


@dataclass
class PendingUser:
    source_file: str
    email: str
    name: str
    microsoft_id: str


def _is_blank(value: str | None) -> bool:
    return value is None or str(value).strip() == ""


def _normalize_header(value: str | None) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _load_object_id_map(object_id_csv: Path) -> dict[str, str]:
    """Map lowercased userPrincipalName -> Microsoft object id."""
    mapping: dict[str, str] = {}
    with object_id_csv.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            upn = (row.get("userPrincipalName") or "").strip()
            obj_id = (row.get("id") or "").strip()
            if _is_blank(upn) or _is_blank(obj_id):
                continue
            mapping[normalize_email(upn)] = obj_id
    return mapping


def _find_header_row(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """Return (header_row_index, column_index_by_key) or None."""
    for row_index, row in enumerate(rows):
        if not row:
            continue
        header_map: dict[str, int] = {}
        for col_index, cell in enumerate(row):
            key = _normalize_header(str(cell) if cell is not None else "")
            if key in (HEADER_NO, HEADER_STUDENT_NAME, HEADER_USERNAME):
                header_map[key] = col_index
        if HEADER_STUDENT_NAME in header_map and HEADER_USERNAME in header_map:
            return row_index, header_map
    return None


def _extract_rows_from_workbook(xlsx_path: Path) -> list[ParsedRow]:
    workbook = load_workbook(xlsx_path, data_only=True)
    parsed: list[ParsedRow] = []
    source_name = xlsx_path.name

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        header_info = _find_header_row(rows)
        if header_info is None:
            continue

        header_row_index, header_map = header_info
        name_col = header_map[HEADER_STUDENT_NAME]
        username_col = header_map[HEADER_USERNAME]

        for row in rows[header_row_index + 1 :]:
            if not row:
                continue
            raw_name = row[name_col] if len(row) > name_col else None
            raw_username = row[username_col] if len(row) > username_col else None
            if _is_blank(raw_name) and _is_blank(raw_username):
                continue
            parsed.append(
                ParsedRow(
                    source_file=source_name,
                    student_name=str(raw_name or "").strip(),
                    username=str(raw_username or "").strip(),
                )
            )
    return parsed


def _write_skip_report_csv(path: Path, rows: list[SkippedRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=SKIP_REPORT_COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "source file": row.source_file,
                    "student name": row.student_name,
                    "username": row.username,
                    "skip reason": row.skip_reason,
                }
            )


def _build_user_instance(*, email: str, name: str, microsoft_id: str, hashed_password: str):
    return UserModel(
        email=email,
        password=hashed_password,
        name=name,
        phone_number=PLACEHOLDER_PHONE,
        communication_email=email,
        roles=[UserModel.UserRole.STUDENT],
        date_of_birth=IMPORT_DATE_OF_BIRTH,
        microsoft_id=microsoft_id,
        is_active=True,
        is_staff=False,
        is_password_change_required=True,
    )


class Command(BaseCommand):
    help = (
        "Import SDEC IGCSE batch workbooks into a tenant (omit --dry-run to write). "
        "Links microsoft_id from object-id CSV by Username/UPN. No course enrollments."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--schema-name",
            type=str,
            default=DEFAULT_SCHEMA_NAME,
            help=f"Tenant schema_name (default {DEFAULT_SCHEMA_NAME!r})",
        )
        parser.add_argument(
            "--xlsx",
            type=str,
            action="append",
            required=True,
            help="Path to an IGCSE batch workbook (repeat for multiple files)",
        )
        parser.add_argument(
            "--object-id-csv",
            type=str,
            required=True,
            help="Path to sdec-user-object-id-list.csv (UTF-8)",
        )
        parser.add_argument(
            "--skip-report-csv",
            type=str,
            default=None,
            help="Path for skipped-row report CSV (default: sdec-ig-batch-skips.csv in cwd)",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and report only; do not write to the database",
        )

    def handle(self, *args, **options):
        schema_name: str = options["schema_name"]
        xlsx_paths = [
            Path(path).expanduser().resolve() for path in options["xlsx"]
        ]
        object_id_path = Path(options["object_id_csv"]).expanduser().resolve()
        skip_report_path = (
            Path(options["skip_report_csv"]).expanduser().resolve()
            if options["skip_report_csv"]
            else Path.cwd() / "sdec-ig-batch-skips.csv"
        )
        dry_run: bool = options["dry_run"]
        commit: bool = not dry_run

        for xlsx_path in xlsx_paths:
            if not xlsx_path.is_file():
                raise CommandError(f"Workbook not found: {xlsx_path}")
        if not object_id_path.is_file():
            raise CommandError(f"Object id CSV not found: {object_id_path}")

        with schema_context(get_public_schema_name()):
            if not Organization.objects.filter(schema_name=schema_name).exists():
                raise CommandError(f"No organization for schema_name={schema_name!r}")

        if dry_run:
            self.stdout.write(
                self.style.WARNING(
                    "DRY RUN — no database writes "
                    "(counts below are planned creates/reuses/backfills)."
                )
            )

        object_id_map = _load_object_id_map(object_id_path)
        parsed_rows: list[ParsedRow] = []
        for xlsx_path in xlsx_paths:
            parsed_rows.extend(_extract_rows_from_workbook(xlsx_path))

        stats = {
            "workbooks": len(xlsx_paths),
            "rows_read": 0,
            "rows_skipped": 0,
            "skipped_missing_name": 0,
            "skipped_missing_username": 0,
            "skipped_duplicate_in_file": 0,
            "skipped_no_object_id": 0,
            "skipped_microsoft_id_conflict": 0,
            "users_created": 0,
            "users_reused": 0,
            "microsoft_ids_backfilled": 0,
        }
        skipped_rows: list[SkippedRow] = []
        pending_users: list[PendingUser] = []
        seen_emails: set[str] = set()

        for row in parsed_rows:
            stats["rows_read"] += 1

            if _is_blank(row.student_name):
                stats["rows_skipped"] += 1
                stats["skipped_missing_name"] += 1
                skipped_rows.append(
                    SkippedRow(
                        source_file=row.source_file,
                        student_name=row.student_name,
                        username=row.username,
                        skip_reason="missing student name",
                    )
                )
                continue

            if _is_blank(row.username):
                stats["rows_skipped"] += 1
                stats["skipped_missing_username"] += 1
                skipped_rows.append(
                    SkippedRow(
                        source_file=row.source_file,
                        student_name=row.student_name,
                        username=row.username,
                        skip_reason="missing username",
                    )
                )
                continue

            email = normalize_email(row.username)
            if email in seen_emails:
                stats["rows_skipped"] += 1
                stats["skipped_duplicate_in_file"] += 1
                skipped_rows.append(
                    SkippedRow(
                        source_file=row.source_file,
                        student_name=row.student_name,
                        username=row.username,
                        skip_reason="duplicate username in import files (first wins)",
                    )
                )
                continue
            seen_emails.add(email)

            microsoft_id = object_id_map.get(email)
            if microsoft_id is None:
                stats["rows_skipped"] += 1
                stats["skipped_no_object_id"] += 1
                skipped_rows.append(
                    SkippedRow(
                        source_file=row.source_file,
                        student_name=row.student_name,
                        username=row.username,
                        skip_reason="no object id for username",
                    )
                )
                continue

            pending_users.append(
                PendingUser(
                    source_file=row.source_file,
                    email=email,
                    name=row.student_name,
                    microsoft_id=microsoft_id,
                )
            )

        with schema_context(schema_name):
            self._apply_users(
                pending_users=pending_users,
                commit=commit,
                stats=stats,
                skipped_rows=skipped_rows,
            )

        self.stdout.write(self.style.SUCCESS(str(stats)))
        _write_skip_report_csv(skip_report_path, skipped_rows)
        self.stdout.write(
            self.style.SUCCESS(
                f"Skip report written to {skip_report_path} ({len(skipped_rows)} rows)"
            )
        )

    def _apply_users(
        self,
        *,
        pending_users: list[PendingUser],
        commit: bool,
        stats: dict,
        skipped_rows: list[SkippedRow],
    ) -> None:
        if not pending_users:
            return

        emails = [item.email for item in pending_users]
        existing_users = {
            user.email.lower(): user
            for user in UserModel.objects.filter(email__in=emails)
        }
        microsoft_id_owners = {
            user.microsoft_id.lower(): user
            for user in UserModel.objects.exclude(microsoft_id__isnull=True)
            .exclude(microsoft_id="")
            .filter(microsoft_id__in=[item.microsoft_id for item in pending_users])
        }

        users_to_create: list[PendingUser] = []
        users_to_backfill: list[tuple[UserModel, PendingUser]] = []

        for item in pending_users:
            owner = microsoft_id_owners.get(item.microsoft_id.lower())
            existing = existing_users.get(item.email)

            if owner is not None and (
                existing is None or owner.pk != existing.pk
            ):
                stats["rows_skipped"] += 1
                stats["skipped_microsoft_id_conflict"] += 1
                skipped_rows.append(
                    SkippedRow(
                        source_file=item.source_file,
                        student_name=item.name,
                        username=item.email,
                        skip_reason=(
                            f"microsoft id already linked to {owner.email}"
                        ),
                    )
                )
                continue

            if existing is None:
                stats["users_created"] += 1
                users_to_create.append(item)
            else:
                stats["users_reused"] += 1
                if _is_blank(existing.microsoft_id):
                    stats["microsoft_ids_backfilled"] += 1
                    users_to_backfill.append((existing, item))

        if not commit:
            return

        with transaction.atomic():
            if users_to_create:
                hashed_password = make_password(IMPORT_PASSWORD)
                UserModel.objects.bulk_create(
                    [
                        _build_user_instance(
                            email=item.email,
                            name=item.name,
                            microsoft_id=item.microsoft_id,
                            hashed_password=hashed_password,
                        )
                        for item in users_to_create
                    ],
                    batch_size=BULK_BATCH_SIZE,
                )

            if users_to_backfill:
                for user, item in users_to_backfill:
                    user.microsoft_id = item.microsoft_id
                UserModel.objects.bulk_update(
                    [user for user, _item in users_to_backfill],
                    ["microsoft_id"],
                    batch_size=BULK_BATCH_SIZE,
                )
