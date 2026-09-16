import io
import unittest
from datetime import date, timedelta
from uuid import uuid4

from django.core.management import call_command
from django.db import connection
from django.test import TestCase, override_settings
from django.utils import timezone
from djmoney.money import Money
from openpyxl import load_workbook
from rest_framework.test import APIClient
from tenant_schemas.utils import get_public_schema_name, schema_context

from app_auth.models import User
from app_course.models import Category, Course, Program, UserCourse
from app_finance.models import PaymentBank, PaymentMethod, UserPayment
from app_finance.payment_receipt_number import ensure_receipt_for_payment
from app_organization.models import Organization
from app_rbac.seeding import seed_rbac


def _database_reachable() -> bool:
    try:
        connection.ensure_connection()
        return True
    except Exception:
        return False


@unittest.skipUnless(_database_reachable(), "PostgreSQL not available")
@override_settings(RBAC_ENFORCE="enforce")
class ExcellentChoiceReportEndpointTests(TestCase):
    schema_name = "xschedjuice"
    api_prefix = "/api/v1"

    @classmethod
    def setUpTestData(cls):
        call_command("migrate_schemas", shared=True, verbosity=0)
        call_command("migrate_schemas", verbosity=0)
        call_command("load-data", schema=cls.schema_name, verbosity=0)

    def setUp(self):
        suffix = uuid4().hex[:6]
        self.today = timezone.localdate()
        self.date_from = self.today.replace(day=1).isoformat()
        self.date_to = self.today.isoformat()
        with schema_context(get_public_schema_name()):
            Organization.objects.filter(schema_name=self.schema_name).update(
                report_style=Organization.ReportStyle.EXCELLENT_CHOICE_STYLE,
            )
        with schema_context(self.schema_name):
            seed_rbac()
            self.finance = User.objects.create_user(
                email=f"ec-fin-{suffix}@example.com",
                password="x",
                name="Finance",
                phone_number="-",
                date_of_birth=date(1990, 1, 1),
                roles=[User.UserRole.FINANCE],
            )
            self.teacher = User.objects.create_user(
                email=f"ec-tch-{suffix}@example.com",
                password="x",
                name="Teacher",
                phone_number="-",
                date_of_birth=date(1990, 1, 1),
                roles=[User.UserRole.TEACHER],
            )
            self.student = User.objects.create_user(
                email=f"ec-stu-{suffix}@example.com",
                password="x",
                name="Student",
                phone_number="-",
                date_of_birth=date(2010, 1, 1),
                roles=[User.UserRole.STUDENT],
            )
            cat = Category.objects.create(name=f"EC API {suffix}")
            prog = Program.objects.create(
                name=f"EC API Prog {suffix}",
                course_creation_method=Program.CourseCreationMethod.MANUAL,
                subject_strategy=Program.SubjectStrategy.NONE,
            )
            self.course = Course.objects.create(
                title=f"EC Course {suffix}",
                category=cat,
                program=prog,
                start_date=self.today,
                end_date=self.today + timedelta(days=30),
            )
            UserCourse.objects.create(
                user=self.student,
                course=self.course,
                assigned_as=UserCourse.AssignedAs.STUDENT,
            )
            self.kpay = PaymentMethod.objects.create(
                name=f"KPay API {suffix}",
                payment_bank=PaymentBank.KPAY,
            )
            payment = UserPayment.objects.create(
                user=self.student,
                course=self.course,
                transaction_id=f"EC-API-{suffix}",
                parsed_amount=Money(120, "USD"),
                actual_amount=Money(120, "USD"),
                payment_method=self.kpay,
                status=UserPayment.Status.PENDING_VERIFICATION,
            )
            payment.status = UserPayment.Status.VERIFIED
            payment.verified_by = self.finance
            payment.payment_date = timezone.now()
            receipt = ensure_receipt_for_payment(payment)
            UserPayment.objects.filter(pk=payment.pk).update(
                status=payment.status,
                verified_by=self.finance,
                payment_date=payment.payment_date,
                receipt=receipt,
            )

    def _client(self, user: User) -> APIClient:
        client = APIClient()
        client.force_authenticate(user=user)
        client.credentials(HTTP_TENANT=self.schema_name)
        return client

    def test_get_returns_rows_columns_and_summary(self):
        resp = self._client(self.finance).get(
            f"{self.api_prefix}/reports/excellent-choice",
            {"date_from": self.date_from, "date_to": self.date_to},
        )
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertFalse(body["isError"])
        self.assertGreaterEqual(len(body["data"]), 1)
        self.assertEqual(body["columns"][0]["key"], "date")
        self.assertEqual(body["data"][0]["cash_received"], "120")
        self.assertEqual(body["summary"]["date"], "SUM")
        self.assertEqual(body["summary"]["voucher_no"], "120")

    def test_teacher_is_forbidden(self):
        resp = self._client(self.teacher).get(
            f"{self.api_prefix}/reports/excellent-choice",
            {"date_from": self.date_from, "date_to": self.date_to},
        )
        self.assertEqual(resp.status_code, 403, resp.content)

    def test_wrong_report_style_is_forbidden(self):
        with schema_context(get_public_schema_name()):
            Organization.objects.filter(schema_name=self.schema_name).update(
                report_style=Organization.ReportStyle.TR_SU_STYLE,
            )
        resp = self._client(self.finance).get(
            f"{self.api_prefix}/reports/excellent-choice",
            {"date_from": self.date_from, "date_to": self.date_to},
        )
        self.assertEqual(resp.status_code, 403, resp.content)

    def test_missing_date_from_is_bad_request(self):
        resp = self._client(self.finance).get(
            f"{self.api_prefix}/reports/excellent-choice",
            {"date_to": self.date_to},
        )
        self.assertEqual(resp.status_code, 400, resp.content)
        self.assertIn("date_from", resp.json()["details"])

    def test_post_returns_xlsx(self):
        with schema_context(get_public_schema_name()):
            Organization.objects.filter(schema_name=self.schema_name).update(
                report_style=Organization.ReportStyle.EXCELLENT_CHOICE_STYLE,
            )
        resp = self._client(self.finance).post(
            f"{self.api_prefix}/reports/excellent-choice",
            {"date_from": self.date_from, "date_to": self.date_to},
            format="json",
        )
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("excellent_choice_", resp["Content-Disposition"])

        workbook = load_workbook(io.BytesIO(resp.content), read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertGreaterEqual(len(rows), 2)
        summary_row = rows[-1]
        self.assertEqual(summary_row[0], "SUM")
        self.assertEqual(summary_row[1], "120")
